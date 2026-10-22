import os
import unittest
from io import BytesIO

from openpyxl import load_workbook
from pycountry import countries

from shipping_excel import build_template


class ShippingAdminExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.previous_env = {key: os.environ.get(key) for key in ("DATABASE_URL", "ADMIN_USERNAME", "ADMIN_PASSWORD", "SECRET_KEY")}
        os.environ["DATABASE_URL"] = "sqlite://"
        os.environ["ADMIN_USERNAME"] = "shipping-test-admin"
        os.environ["ADMIN_PASSWORD"] = "shipping-test-password"
        os.environ["SECRET_KEY"] = "shipping-test-secret"
        from app import create_app
        from models import db
        cls.db = db
        cls.app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite://",
                              "SECRET_KEY": "shipping-test-secret"})

    def setUp(self):
        with self.app.app_context():
            self.db.drop_all()
            self.db.create_all()
        self.client = self.app.test_client()
        login = self.client.post("/api/admin/login", json={"username": "shipping-test-admin", "password": "shipping-test-password"})
        self.assertEqual(login.status_code, 200)
        self.headers = {"Authorization": "Bearer " + login.get_json()["token"]}
        self.client.put("/api/admin/shipping/mode", json={"mode": "manual_excel"}, headers=self.headers)

    def workbook(self, dhl=None, fedex=None, community=None):
        output = build_template(list(countries))
        workbook = load_workbook(output)
        sheet = workbook["Shipping Rates"]
        row = next(index for index in range(2, sheet.max_row + 1) if sheet.cell(index, 2).value == "GM")
        for column, value in ((4, dhl), (5, fedex), (6, community)):
            if value is not None:
                sheet.cell(row, column).value = value
        result = BytesIO()
        workbook.save(result)
        return result.getvalue()

    def preview(self, data):
        return self.client.post("/api/admin/shipping/excel-preview", headers=self.headers,
            data={"file": (BytesIO(data), "rates.xlsx")}, content_type="multipart/form-data")

    def confirm(self, import_id, approve=True, confirm_zero=False):
        return self.client.post("/api/admin/shipping/excel-confirm", headers=self.headers, json={
            "import_id": import_id, "confirm_import": True, "approve_prices": approve,
            "confirm_zero_prices": confirm_zero})

    def test_template_download_contains_every_country_and_required_columns(self):
        response = self.client.get("/api/admin/shipping/excel-template", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        sheet = workbook["Shipping Rates"]
        self.assertEqual(sheet.max_row - 1, len(list(countries)))
        self.assertEqual([cell.value for cell in sheet[1]], [
            "Country Name", "ISO Country Code", "Origin", "DHL Price USD per 0.5 kg",
            "FedEx Price USD per 0.5 kg", "Community Courier Price USD per 0.5 kg", "Optional Notes"])

    def test_import_updates_canonical_rates_and_blank_cells_preserve_existing_prices(self):
        from models import Category, Product, ShippingRate, ShippingRateHistory
        from shipping_core import shipping_rate_is_current
        with self.app.app_context():
            old = ShippingRate(country_code="GM", country_name="Gambia", provider="dhl", base_price_usd=3,
                               status="available", approved=True, manual_override=True, source_type="manual_admin")
            old_fedex = ShippingRate(country_code="GM", country_name="Gambia", provider="fedex", base_price_usd=2,
                                     status="available", approved=True, manual_override=True, source_type="manual_admin")
            self.db.session.add_all([old, old_fedex])
            self.db.session.commit()
        preview = self.preview(self.workbook(fedex=4.25, community=0))
        self.assertEqual(preview.status_code, 201)
        summary = preview.get_json()["preview"]
        self.assertEqual(summary["invalid_rows"], 0)
        self.assertEqual(summary["valid_prices"], 2)
        self.assertEqual(summary["existing_prices_to_update"], 1)
        self.assertEqual(summary["prices_to_create"], 1)
        self.assertGreater(summary["blank_prices_unchanged"], 0)
        self.assertEqual(self.confirm(preview.get_json()["import_id"], confirm_zero=True).status_code, 200)
        with self.app.app_context():
            rates = {rate.provider: rate for rate in ShippingRate.query.filter_by(country_code="GM").all()}
            self.assertEqual(str(rates["dhl"].base_price_usd), "3.00")
            self.assertEqual(str(rates["fedex"].base_price_usd), "4.25")
            self.assertEqual(str(rates["community"].base_price_usd), "0.00")
            self.assertTrue(all(rate.approved and rate.source_type == "excel_import" for rate in (rates["fedex"], rates["community"])))
            self.assertTrue(shipping_rate_is_current(rates["community"]))
            self.assertEqual(ShippingRateHistory.query.filter_by(rate_id=rates["dhl"].id).count(), 0)
            self.assertEqual(ShippingRateHistory.query.filter_by(rate_id=rates["fedex"].id).count(), 1)
            category = Category(name="Test", slug="test")
            self.db.session.add(category)
            self.db.session.flush()
            product = Product(name="Test product", slug="test-product", description="", price=1,
                              stock=1, weight_kg=0.3, category_id=category.id)
            self.db.session.add(product)
            self.db.session.commit()
            product_id = product.id
        from unittest.mock import patch
        with patch("shipping.groq_research_destination", side_effect=AssertionError("checkout must not call AI")):
            quote = self.client.post("/api/shipping/estimate", json={"country_code": "GM", "provider": "fedex",
                "items": [{"product_id": product_id, "quantity": 2}]})
        self.assertEqual(quote.get_json()["status"], "available")
        self.assertEqual(quote.get_json()["price_usd"], "8.50")

    def test_invalid_rows_block_the_entire_import_and_mode_blocks_upload(self):
        from models import ShippingRate
        preview = self.preview(self.workbook(fedex=-3))
        self.assertEqual(preview.status_code, 201)
        self.assertGreater(preview.get_json()["preview"]["invalid_rows"], 0)
        self.assertIsNone(preview.get_json()["import_id"])
        with self.app.app_context():
            self.assertEqual(ShippingRate.query.count(), 0)
        self.client.put("/api/admin/shipping/mode", json={"mode": "ai_research"}, headers=self.headers)
        response = self.preview(self.workbook(fedex=3))
        self.assertEqual(response.status_code, 409)

    def test_excel_rates_require_approval_and_zero_confirmation(self):
        from models import ShippingRate
        preview = self.preview(self.workbook(community=0))
        import_id = preview.get_json()["import_id"]
        self.assertEqual(self.confirm(import_id, confirm_zero=False).status_code, 400)
        self.assertEqual(self.confirm(import_id, approve=False, confirm_zero=True).status_code, 200)
        with self.app.app_context():
            rate = ShippingRate.query.filter_by(country_code="GM", provider="community").one()
            self.assertEqual(rate.status, "manual_review")
            self.assertFalse(rate.approved)
            self.assertEqual(rate.source_type, "excel_import")
            self.assertTrue(rate.manual_override)

    @classmethod
    def tearDownClass(cls):
        cls.db.session.remove()
        cls.db.engines.clear()
        for key, value in cls.previous_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


if __name__ == "__main__":
    unittest.main()
