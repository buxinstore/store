import unittest
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from shipping_excel import HEADERS, ORIGIN, build_template, parse_workbook


COUNTRIES = [SimpleNamespace(name="Gambia", alpha_2="GM"), SimpleNamespace(name="United States", alpha_2="US")]


def workbook_bytes(country_items=COUNTRIES):
    file = build_template(country_items)
    return file.getvalue()


class ShippingExcelTests(unittest.TestCase):
    def test_template_has_complete_country_rows_and_instruction_sheet(self):
        data = workbook_bytes()
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Shipping Rates", "Instructions"])
        sheet = workbook["Shipping Rates"]
        self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
        self.assertEqual(sheet.max_row, len(COUNTRIES) + 1)
        self.assertEqual([sheet.cell(row, 2).value for row in (2, 3)], ["GM", "US"])
        self.assertEqual(sheet["C2"].value, ORIGIN)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)
        self.assertFalse(sheet["D2"].protection.locked)
        self.assertTrue(sheet["A2"].protection.locked)
        self.assertIn("Blank price cells", workbook["Instructions"]["A3"].value)

    def test_valid_prices_for_all_providers_and_blank_cells(self):
        workbook = load_workbook(BytesIO(workbook_bytes()))
        sheet = workbook["Shipping Rates"]
        sheet["D2"] = 4.5
        sheet["E2"] = 7
        sheet["F2"] = 0
        output = BytesIO()
        workbook.save(output)
        parsed = parse_workbook(output.getvalue(), COUNTRIES)
        self.assertEqual(parsed["invalid_rows"], 0)
        self.assertEqual(parsed["rows"][0]["prices"], {"dhl": "4.50", "fedex": "7.00", "community": "0.00"})
        self.assertEqual(parsed["zeros"], [{"row": 2, "code": "GM", "provider": "community"}])
        self.assertEqual(parsed["rows"][1]["prices"], {})

    def test_duplicate_codes_and_bad_country_identity_are_errors(self):
        workbook = load_workbook(BytesIO(workbook_bytes()))
        sheet = workbook["Shipping Rates"]
        sheet["B3"] = "GM"
        output = BytesIO()
        workbook.save(output)
        parsed = parse_workbook(output.getvalue(), COUNTRIES)
        self.assertEqual(parsed["invalid_rows"], 2)
        self.assertTrue(any("Duplicate ISO" in item["message"] for item in parsed["errors"]))

    def test_formulas_and_text_prices_are_rejected(self):
        workbook = load_workbook(BytesIO(workbook_bytes()))
        sheet = workbook["Shipping Rates"]
        sheet["D2"] = "=1+1"
        sheet["E2"] = "12.50"
        output = BytesIO()
        workbook.save(output)
        parsed = parse_workbook(output.getvalue(), COUNTRIES)
        self.assertEqual(parsed["invalid_rows"], 1)
        self.assertEqual(len(parsed["errors"]), 2)
        self.assertEqual([item["country_code"] for item in parsed["rows"]], ["US"])

    def test_malformed_and_oversized_workbooks_are_rejected(self):
        with self.assertRaises(ValueError):
            parse_workbook(b"not an xlsx", COUNTRIES)
        with self.assertRaises(ValueError):
            parse_workbook(b"x" * (2 * 1024 * 1024 + 1), COUNTRIES)


if __name__ == "__main__":
    unittest.main()
