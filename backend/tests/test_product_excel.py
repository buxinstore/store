import unittest
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from product_excel import build_product_template, parse_product_workbook


class ProductExcelTests(unittest.TestCase):
    def setUp(self):
        self.category = SimpleNamespace(id=7, name="Sensors", slug="sensors")

    def workbook(self, *, sku="SENSOR-1", product_name="Temperature Sensor", price=4.25, stock=12,
                 category="Sensors", primary_url=None, media=None, second_product=None):
        output = build_product_template()
        workbook = load_workbook(output)
        sheet = workbook["Products"]
        row = [product_name, sku, "", "A short description", "A longer description", category,
               "", "Maker", "sensor, temperature", price, None, "USD", stock, "TRUE", "Active",
               "TRUE", 0.25, "kg", "TRUE", "MFR-1", "TS-1", "Range: -40 to 80 C",
               "https://example.com/datasheet.pdf", "One year", "GM", primary_url]
        sheet.append(row)
        if second_product:
            sheet.append(second_product)
        if media:
            media_sheet = workbook["Images and Videos"]
            for record in media:
                media_sheet.append(record)
        result = BytesIO()
        workbook.save(result)
        return result.getvalue()

    def parse(self, content, existing=None):
        return parse_product_workbook(content, [self.category], existing or {})

    def test_downloaded_template_has_only_importable_sheets_and_headers(self):
        workbook = load_workbook(build_product_template(), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Instructions", "Products", "Images and Videos"])
        headers = [cell.value for cell in next(workbook["Products"].iter_rows(min_row=1, max_row=1))]
        self.assertIn("Selling Price", headers)
        self.assertNotIn("Cost Price", headers)
        workbook.close()

    def test_product_with_multiple_images_and_video_is_parsed(self):
        media = [
            ["SENSOR-1", "Image", "https://images.example.com/a.jpg", "Front", "TRUE", 0],
            ["SENSOR-1", "Image", "https://images.example.com/b.jpg", "Back", "FALSE", 1],
            ["SENSOR-1", "Video", "https://youtu.be/abcdefghijk", "Demo", "FALSE", 2],
        ]
        result = self.parse(self.workbook(media=media))
        self.assertEqual(result["errors"], [])
        self.assertEqual(len(result["products"]), 1)
        self.assertEqual([item["type"] for item in result["media"]], ["image", "image", "video"])
        self.assertEqual(result["products"][0]["specifications"], {"Range": "-40 to 80 C"})

    def test_missing_optional_images_and_videos_are_valid(self):
        result = self.parse(self.workbook())
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["media"], [])

    def test_duplicate_skus_and_invalid_category_are_reported(self):
        second = ["Another Sensor", "sensor-1", "another-sensor", "", "", "Unknown", "", "", "", 1, None,
                  "USD", 2, "TRUE", "Active", "FALSE", 0.1, "kg", "TRUE", "", "", "", "", "", "", ""]
        result = self.parse(self.workbook(second_product=second))
        messages = " ".join(error["message"] for error in result["errors"])
        self.assertIn("Duplicate SKU", messages)
        self.assertIn("does not match an existing category", messages)

    def test_existing_sku_is_reported_as_update_and_blank_fields_remain_unprovided(self):
        result = self.parse(self.workbook(), {"sensor-1": object()})
        row = result["products"][0]
        self.assertEqual(result["update_count"], 1)
        self.assertNotIn("Original Price", row["provided"])
        self.assertNotIn("Subcategory", row["provided"])

    def test_invalid_and_private_media_urls_are_rejected(self):
        media = [["SENSOR-1", "Video", "https://example.com/video.mp4", "Not an approved host", "FALSE", 0],
                 ["UNKNOWN", "Image", "https://127.0.0.1/secret.png", "Private", "FALSE", 0]]
        result = self.parse(self.workbook(media=media))
        self.assertEqual(len(result["media"]), 0)
        self.assertGreaterEqual(len(result["errors"]), 2)

    def test_formula_cells_are_rejected(self):
        content = self.workbook()
        workbook = load_workbook(BytesIO(content))
        workbook["Products"]["D3"] = "=HYPERLINK(\"https://example.com\",\"click\")"
        result = BytesIO()
        workbook.save(result)
        with self.assertRaisesRegex(ValueError, "Formula cells are not allowed"):
            self.parse(result.getvalue())

    def test_malformed_workbook_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Could not read"):
            self.parse(b"not an xlsx")


if __name__ == "__main__":
    unittest.main()
