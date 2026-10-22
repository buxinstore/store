"""Validated XLSX template and parsing for canonical shipping-rate imports."""
from io import BytesIO
from copy import copy
import zipfile
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = [
    "Country Name", "ISO Country Code", "Origin", "DHL Price USD per 0.5 kg",
    "FedEx Price USD per 0.5 kg", "Community Courier Price USD per 0.5 kg", "Optional Notes",
]
PRICE_COLUMNS = {"DHL Price USD per 0.5 kg": "dhl", "FedEx Price USD per 0.5 kg": "fedex",
                 "Community Courier Price USD per 0.5 kg": "community"}
ORIGIN = "Shenzhen, China"
MAX_XLSX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 12 * 1024 * 1024
MAX_DATA_ROWS = 600


def build_template(country_items):
    """Return a complete country template as an in-memory XLSX file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Shipping Rates"
    sheet.append(HEADERS)
    for country in country_items:
        sheet.append([country.name, country.alpha_2, ORIGIN, None, None, None, None])

    header_fill = PatternFill("solid", fgColor="17365D")
    identity_fill = PatternFill("solid", fgColor="D9EAF7")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row[:3]:
            cell.fill = identity_fill
            protection = copy(cell.protection)
            protection.locked = True
            cell.protection = protection
        for cell in row[3:6]:
            cell.fill = input_fill
            cell.number_format = '"$"#,##0.00'
            protection = copy(cell.protection)
            protection.locked = False
            cell.protection = protection
        row[6].fill = input_fill
        protection = copy(row[6].protection)
        protection.locked = False
        row[6].protection = protection
    price_validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000", allow_blank=True)
    price_validation.error = "Enter a numeric USD price from 0 to 100000. A zero price requires explicit confirmation during import."
    price_validation.errorTitle = "Invalid shipping price"
    price_validation.prompt = "Base price per first 0.5 kg, in USD."
    price_validation.promptTitle = "Shipping price"
    sheet.add_data_validation(price_validation)
    price_validation.add(f"D2:F{len(country_items) + 1}")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{len(country_items) + 1}"
    sheet.row_dimensions[1].height = 34
    for column, width in {"A": 30, "B": 20, "C": 24, "D": 27, "E": 29, "F": 42, "G": 42}.items():
        sheet.column_dimensions[column].width = width
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Buxin Store Shipping Rate Import Instructions"])
    instructions.append(["1. Fill yellow price cells with numeric USD amounts per 0.5 kg (the base price for the first 0.5 kg)."])
    instructions.append(["2. Blank price cells leave existing saved rates unchanged."])
    instructions.append(["3. Leave a price blank when it is unverified. Never invent or guess a price."])
    instructions.append(["4. Do not change country names, ISO country codes, or the Shenzhen, China origin."])
    instructions.append(["5. Review the Admin preview. Confirm approval only for rates you authorize for checkout."])
    instructions.append(["6. A numeric zero is treated as a real price and requires a separate explicit confirmation."])
    instructions.append(["7. Upload this .xlsx file in Admin → Shipping Settings → Manual Excel Mode."])
    instructions.append(["8. Importing adds or updates only nonblank provider prices; it never deletes a price."])
    instructions.column_dimensions["A"].width = 115
    instructions.freeze_panes = "A2"
    for cell in instructions[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = header_fill
    for row in instructions.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    result = BytesIO()
    workbook.save(result)
    result.seek(0)
    return result


def parse_workbook(data, country_items):
    """Validate a submitted workbook completely before returning importable rows."""
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ValueError("The Excel file is empty or exceeds the 2 MB upload limit.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("The workbook expands beyond the safe 12 MB limit.")
            if "xl/vbaProject.bin" in archive.namelist():
                raise ValueError("Macro-enabled workbooks are not supported.")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("This is not a valid .xlsx workbook.") from exc
    required_sheets = {"Shipping Rates", "Instructions"}
    if set(workbook.sheetnames) != required_sheets:
        raise ValueError("The workbook must contain exactly the Shipping Rates and Instructions worksheets.")
    sheet = workbook["Shipping Rates"]
    if sheet.max_row > MAX_DATA_ROWS + 1 or sheet.max_column > 20:
        raise ValueError("The Shipping Rates worksheet contains too many rows or columns.")
    header = [sheet.cell(1, index).value for index in range(1, len(HEADERS) + 1)]
    if header != HEADERS or sheet.max_column != len(HEADERS):
        raise ValueError("The Shipping Rates header row does not match the required template columns.")
    countries = {country.alpha_2: country for country in country_items}
    seen = {}
    parsed, errors, zeros = [], [], []
    for row_number in range(2, sheet.max_row + 1):
        cells = [sheet.cell(row_number, column) for column in range(1, 8)]
        values = [cell.value for cell in cells]
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        row_errors = []
        name = str(values[0] or "").strip()
        code = str(values[1] or "").strip()
        origin = str(values[2] or "").strip()
        country = countries.get(code)
        if not country:
            row_errors.append("ISO country code is not in the supported destination list.")
        else:
            if code != code.upper():
                row_errors.append("ISO country code must use uppercase letters and must not be changed.")
            if name != country.name:
                row_errors.append(f"Country name must exactly match {country.name} for ISO code {code}.")
        if origin != ORIGIN:
            row_errors.append(f"Origin must remain {ORIGIN}.")
        if code in seen:
            row_errors.append(f"Duplicate ISO country code; it also appears on row {seen[code]}.")
            errors.append({"row": seen[code], "code": code, "message": f"Duplicate ISO country code also appears on row {row_number}."})
        else:
            seen[code] = row_number

        prices = {}
        for index, header_name in enumerate(HEADERS[3:6], start=3):
            value = values[index]
            if cells[index].data_type == "f":
                row_errors.append(f"{header_name} cannot contain a formula.")
                continue
            if value is None or (isinstance(value, str) and value.strip() == ""):
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                row_errors.append(f"{header_name} must be a numeric cell or blank.")
                continue
            try:
                amount = Decimal(str(value))
                rounded = amount.quantize(Decimal("0.01"))
            except (InvalidOperation, ValueError):
                row_errors.append(f"{header_name} must be a finite numeric price.")
                continue
            if not amount.is_finite() or amount != rounded:
                row_errors.append(f"{header_name} must be finite and have no more than two decimal places.")
                continue
            if amount < 0 or amount > 100000:
                row_errors.append(f"{header_name} must be between USD 0 and USD 100000.")
                continue
            prices[PRICE_COLUMNS[header_name]] = f"{rounded:.2f}"
            if amount == 0:
                zeros.append({"row": row_number, "code": code, "provider": PRICE_COLUMNS[header_name]})
        if row_errors:
            errors.extend({"row": row_number, "code": code, "message": message} for message in row_errors)
        else:
            parsed.append({"row": row_number, "country_code": code, "country_name": country.name,
                           "origin": ORIGIN, "prices": prices,
                           "notes": str(values[6] or "").strip()[:1000]})
    workbook.close()
    if not parsed and not errors:
        errors.append({"row": 2, "code": "", "message": "No destination rows were found. Use the complete downloaded template."})
    return {"rows": parsed, "errors": errors, "zeros": zeros, "recognized_rows": len(parsed),
            "rows_seen": len(seen), "invalid_rows": len({item["row"] for item in errors})}
