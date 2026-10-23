"""Simple product workbook template and validation for Admin → Products."""
from io import BytesIO
import ipaddress
import math
import re
import zipfile
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PRODUCT_HEADERS = ["Product name", "SKU", "Category", "Price (USD)", "Stock", "Product image"]
MAX_XLSX_BYTES = 4 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 24 * 1024 * 1024
MAX_PRODUCTS = 1500
ALLOWED_VIDEO_HOSTS = {"youtube.com", "www.youtube.com", "youtu.be", "vimeo.com", "www.vimeo.com", "player.vimeo.com"}


def safe_image_url(value):
    value = str(value or "").strip()
    if not value:
        return ""
    try:
        parsed = urlsplit(value)
        host = (parsed.hostname or "").lower().rstrip(".")
        if parsed.scheme != "https" or not host or parsed.username or parsed.password or parsed.port not in (None, 443):
            return ""
        try:
            address = ipaddress.ip_address(host)
            if not address.is_global:
                return ""
        except ValueError:
            if host == "localhost" or host.endswith((".localhost", ".local", ".internal")):
                return ""
        return value[:2000]
    except ValueError:
        return ""


def safe_external_url(value, kind="image"):
    """Keep the product-media API's existing validator interface available."""
    url = safe_image_url(value)
    if not url or kind != "video":
        return url if kind == "image" else ""
    host = (urlsplit(url).hostname or "").lower().rstrip(".")
    return url if host in ALLOWED_VIDEO_HOSTS else ""


def _slug(value):
    return re.sub(r"[^a-z0-9]+", "-", str(value or "").lower()).strip("-")[:200]


def build_product_template():
    workbook = Workbook()
    products = workbook.active
    products.title = "Products"
    products.append(PRODUCT_HEADERS)
    products.freeze_panes = "A2"
    products.auto_filter.ref = "A1:F1"
    widths = [28, 20, 24, 16, 12, 44]
    for index, (header, width) in enumerate(zip(PRODUCT_HEADERS, widths), start=1):
        cell = products.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        products.column_dimensions[cell.column_letter].width = width
    products.row_dimensions[1].height = 30
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def parse_product_workbook(data, categories, existing_by_sku):
    if not isinstance(data, bytes) or not data or len(data) > MAX_XLSX_BYTES:
        raise ValueError("Workbook must be a non-empty .xlsx file no larger than 4 MB")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if len(archive.namelist()) > 100 or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Workbook expands beyond the allowed size")
            if any(item.file_size and item.compress_size == 0 for item in archive.infolist()):
                raise ValueError("Workbook contains an invalid compressed entry")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Could not read the .xlsx workbook") from exc
    if "Products" not in workbook.sheetnames or set(workbook.sheetnames) - {"Products", "Categories"}:
        workbook.close()
        raise ValueError("Workbook must contain a Products sheet and may include its Categories sheet")

    existing = {str(key).strip().casefold(): value for key, value in (existing_by_sku or {}).items()}
    category_map = {}
    for category in categories:
        category_map[str(category.name).strip().casefold()] = category
        category_map[str(category.slug).strip().casefold()] = category
    sheet = workbook["Products"]
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                workbook.close()
                raise ValueError(f"Formula cells are not allowed ({sheet.title}!{cell.coordinate})")
    headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    extra_headers = headers[len(PRODUCT_HEADERS):]
    if headers[:len(PRODUCT_HEADERS)] != PRODUCT_HEADERS or (extra_headers and extra_headers != ["", "How to use"]):
        workbook.close()
        raise ValueError("Products sheet must have the six product columns; optional legacy columns are blank and How to use")
    if "Categories" in workbook.sheetnames:
        category_sheet = workbook["Categories"]
        if str(category_sheet.cell(1, 1).value or "").strip().casefold() != "category":
            workbook.close()
            raise ValueError("Categories sheet must have a Category header in its first cell")
    if sheet.max_row > MAX_PRODUCTS + 1:
        workbook.close()
        raise ValueError("Workbook contains more than 1,500 product rows")

    products, errors, seen = [], [], set()
    detected = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values or ())
        row = {header: values[index] if index < len(values) else None for index, header in enumerate(PRODUCT_HEADERS)}
        if not any(value not in (None, "") for value in row.values()):
            continue
        detected += 1
        row_errors = []
        name = str(row["Product name"] or "").strip()
        sku = str(row["SKU"] or "").strip().upper()
        category_name = str(row["Category"] or "").strip()
        category = category_map.get(category_name.casefold())
        if not name:
            row_errors.append("Product name is required")
        elif len(name) > 180:
            row_errors.append("Product name must be 180 characters or fewer")
        if not sku:
            row_errors.append("SKU is required")
        elif len(sku) > 100:
            row_errors.append("SKU must be 100 characters or fewer")
        if sku and sku.casefold() in seen:
            row_errors.append(f"Duplicate SKU '{sku}' in workbook")
        if sku:
            seen.add(sku.casefold())
        if not category:
            row_errors.append(f"Category '{category_name or '(blank)'}' does not match an existing category")
        try:
            raw_price = row["Price (USD)"]
            price_value = ("" if raw_price is None else str(raw_price)).strip().replace("$", "").replace(",", "")
            price = float(price_value)
            if not math.isfinite(price) or price < 0 or price > 99999999.99:
                raise ValueError
        except (TypeError, ValueError, OverflowError):
            row_errors.append("Price (USD) must be a non-negative amount")
            price = 0
        try:
            raw_stock = float(row["Stock"])
            if not math.isfinite(raw_stock) or raw_stock < 0 or int(raw_stock) != raw_stock or raw_stock > 2147483647:
                raise ValueError
            stock = int(raw_stock)
        except (TypeError, ValueError, OverflowError):
            row_errors.append("Stock must be a non-negative whole number")
            stock = 0
        image_url = safe_image_url(row["Product image"])
        if row["Product image"] not in (None, "") and not image_url:
            row_errors.append("Product image must be a public HTTPS image URL")
        if row_errors:
            errors.extend({"row": row_number, "sheet": "Products", "message": message} for message in row_errors)
            continue
        products.append({
            "row": row_number,
            "provided": [key for key, value in row.items() if value not in (None, "")],
            "name": name,
            "sku": sku,
            "slug": _slug(name),
            "category_id": category.id,
            "category_name": category.name,
            "price": price_value,
            "stock": stock,
            "primary_image_url": image_url,
            "short_description": "",
            "description": "",
            "subcategory": "",
            "brand": "",
            "compare_at_price": None,
            "weight_kg": None,
            "status": "active",
            "is_featured": False,
            "details": {},
            "specifications": None,
            "tags": None,
        })
    workbook.close()
    if detected == 0:
        raise ValueError("Products worksheet does not contain any product rows")
    return {
        "products": products,
        "media": [],
        "errors": errors,
        "rows_detected": detected,
        "new_count": sum(1 for product in products if product["sku"].casefold() not in existing),
        "update_count": sum(1 for product in products if product["sku"].casefold() in existing),
    }
